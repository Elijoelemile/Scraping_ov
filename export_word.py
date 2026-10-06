"""Export Word / PDF du comparatif affiché dans l'application (1 à N sites)."""
import io
import os
import shutil
import subprocess
import tempfile
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm

from rapport_ovoyages import MOIS_FR, BLEU_FONCE, GRIS, ombrer, ecrire, ecrire_ligne

VERT = "C6EFCE"
AUCUNE = "Aucun départ"  # case de prix sans départ
UN_SEUL = "Un seul site"  # écart impossible : un seul site a un prix
NON_RELEVE = "Non relevé"  # le site n'a pas pu fournir ce mois (coupure, panne)
INCOMPLET = "Incomplet"  # écart impossible : un site n'a pas pu être relevé ce mois-là


def _figer(table, largeurs):
    """Lignes insécables, tableau d'un seul tenant, largeurs de colonnes."""
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, w in zip(row.cells, largeurs):
            c.width = w
            if i < len(table.rows) - 1:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True


def construire_docx(produit, sites, contexte, large, synthese, releve_le, paire=None):
    """large : DataFrame (une ligne par date) avec colonnes mois, semaine, jour, date, <site>…, moins_cher, ecart,
    et <site>__mp (meilleur prix). synthese : DataFrame par site.
    paire : 2 sites choisis ; seul l'« Écart » porte alors sur ces deux sites (« Moins cher » et le surlignage
    restent le moins cher de la ligne, parmi tous les sites)."""
    suffixe = f" ({paire[0]} / {paire[1]})" if paire else ""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Comparatif de prix" if len(sites) > 1 else "Relevé de prix", level=0)
    p = doc.add_paragraph()
    p.add_run(f"Produit : {produit}").bold = True
    for ligne in contexte:
        p.add_run("\n" + ligne)
    p.add_run(f"\nPrix relevés le {releve_le:%d/%m/%Y} à {releve_le:%H:%M}. Prix par personne (base chambre double) ; "
              "les tarifs évoluent en continu.")
    if len(sites) > 1:
        doc.add_paragraph().add_run(
            "Le prix le plus bas de chaque date est surligné en vert. "
            + (f"Écart = différence de prix entre {paire[0]} et {paire[1]}. " if paire else
               "Écart = prix le plus haut − prix le plus bas parmi les sites. ")
            + "Une semaine correspond à une ligne du calendrier (du lundi au dimanche).").italic = True

    # Synthèse
    doc.add_heading("Synthèse", level=1)
    cols = list(synthese.columns)
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, cols):
        ecrire(c, str(h), gras=True, couleur="FFFFFF", taille=9)
        ombrer(c, BLEU_FONCE)
    for _, r in synthese.iterrows():
        cells = t.add_row().cells
        for c, h in zip(cells, cols):
            ecrire(c, str(r[h]), gras=h == cols[0], taille=9)
    _figer(t, [Cm(17.8 / len(cols))] * len(cols))

    # Un tableau par mois
    entetes = ["Semaine", "Jour", "Date"] + sites + (["Moins cher", f"Écart{suffixe}"] if len(sites) > 1 else [])
    reste = 17.8 - 2.2 - 1.6 - 1.6 - (2.6 + 1.8 if len(sites) > 1 else 0)
    largeurs = [Cm(2.2), Cm(1.6), Cm(1.6)] + [Cm(reste / len(sites))] * len(sites) + \
               ([Cm(2.6), Cm(1.8)] if len(sites) > 1 else [])
    for (aaaa, mm), bloc in large.groupby(["aaaa", "mm"], sort=True):  # ordre chronologique
        doc.add_paragraph()
        h = doc.add_heading(f"{MOIS_FR[mm]} {aaaa}", level=1)
        h.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=len(entetes))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, txt in zip(t.rows[0].cells, entetes):
            ecrire(c, txt, gras=True, couleur="FFFFFF", taille=9)
            ombrer(c, BLEU_FONCE)
        precedente, cellule_semaine = None, None
        for _, r in bloc.sort_values("date").iterrows():
            cells = t.add_row().cells
            if r["semaine"] == precedente:
                cellule_semaine = cellule_semaine.merge(cells[0])  # « Semaine N » fusionnée sur ses jours
            else:
                cellule_semaine = cells[0]
                ecrire(cellule_semaine, f"Semaine {r['semaine']}", gras=True)
                precedente = r["semaine"]
            ecrire(cells[1], r["jour"])
            ecrire(cells[2], r["date"].strftime("%d/%m"))
            prix = [r[s] for s in sites]
            mini = min((x for x in prix if x == x), default=None)  # x == x : écarte les NaN
            for c, s, x in zip(cells[3:], sites, prix):
                if x != x:
                    manquant = r.get(f"{s}__echec") == True  # noqa: E712
                    ecrire(c, NON_RELEVE if manquant else AUCUNE, couleur="B45309" if manquant else "9AA5B1", taille=8)
                    continue
                ecrire(c, f"{int(x)} €", gras=True, taille=10)
                if len(sites) > 1 and x == mini:
                    ombrer(c, VERT)
                if r.get(f"{s}__mp") == True:  # noqa: E712 (la valeur peut être NaN, qui est « vraie »)
                    ecrire_ligne(c, "Meilleur prix du site", taille=7)
            if len(sites) > 1:
                ecrire(cells[-2], r["moins_cher"] or AUCUNE, gras=bool(r["moins_cher"]), taille=9)
                if r["ecart"] == r["ecart"]:  # r["ecart"] != r["ecart"] signifie NaN
                    ecrire(cells[-1], f"{int(r['ecart'])} €", taille=9)
                else:
                    aucun = paire and all(r[s] != r[s] for s in paire)  # aucun des deux sites n'a de départ
                    ecrire(cells[-1], INCOMPLET if r.get("incomplet_ecart", r.get("incomplet")) == True  # noqa: E712
                           else AUCUNE if aucun else UN_SEUL, couleur="9AA5B1", taille=8)
        _figer(t, largeurs)

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()


def docx_en_pdf(docx_bytes):
    """Conversion en PDF : Microsoft Word sous Windows, sinon LibreOffice (serveurs Linux, Streamlit Cloud).
    Renvoie None si aucun des deux n'est disponible."""
    with tempfile.TemporaryDirectory() as d:
        src, dst = os.path.join(d, "rapport.docx"), os.path.join(d, "rapport.pdf")
        with open(src, "wb") as f:
            f.write(docx_bytes)
        if os.name == "nt":
            script = (f"$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
                      f"$d=$w.Documents.Open('{src}',$false,$true); $d.ExportAsFixedFormat('{dst}',17); "
                      f"$d.Close($false); $w.Quit()")
            try:
                subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, timeout=180,
                               capture_output=True)
            except Exception:
                pass
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not os.path.exists(dst) and soffice:
            try:
                # profil LibreOffice dans le dossier temporaire : évite les conflits entre conversions
                subprocess.run([soffice, f"-env:UserInstallation=file://{d}/profil", "--headless",
                                "--convert-to", "pdf", "--outdir", d, src],
                               check=True, timeout=180, capture_output=True)
            except Exception:
                pass
        if os.path.exists(dst):
            with open(dst, "rb") as f:
                return f.read()
        return None
