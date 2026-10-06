"""Export Excel du comparatif affiché dans l'application (1 à N sites).

Trois onglets :
  Synthèse   — contexte du relevé + tableau par site (formules sur l'onglet Comparatif)
  Comparatif — une ligne par date, une colonne de prix par site ; « Moins cher » et « Écart » en formules
  Détail     — les lignes brutes (site, date, prix, meilleur prix, compagnie)
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rapport_ovoyages import MOIS_FR

POLICE = "Arial"
BLEU_FONCE = "2E3E50"
VERT = "C6EFCE"
GRIS = "EAEEF2"
FORMAT_EUR = '#,##0 "€"'
FORMAT_DATE = "DD/MM/YYYY"
AUCUNE = "Aucun départ"  # case de prix sans départ (texte : ignoré par COUNT, MIN, AVERAGE)
UN_SEUL = "Un seul site"  # écart impossible : un seul site a un prix
NON_RELEVE = "Non relevé"  # le site n'a pas pu fournir ce mois (coupure, panne)
INCOMPLET = "Incomplet"  # un site n'a pas pu être relevé ce mois-là
NON_PRECISEE = "Non précisée"
EXTRAS_DETAIL = [("retour", "Date de retour"), ("formule", "Formule"), ("vol_direct", "Vol direct"),
                 ("voyagiste", "Voyagiste")]  # compagnie aérienne absente (le site ne l'indique pas)


def _entete(ws, ligne, titres):
    for j, t in enumerate(titres, start=1):
        c = ws.cell(ligne, j, t)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU_FONCE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ecrire(ws, ligne, col, valeur, fmt=None, gras=False, centre=False):
    c = ws.cell(ligne, col, valeur)
    c.font = Font(name=POLICE, bold=gras)
    if fmt:
        c.number_format = fmt
    if centre:
        c.alignment = Alignment(horizontal="center")
    return c


def formule_moins_cher(plage, l0, lf):
    """Nom du site le moins cher d'une ligne parmi tous les sites (ou « Égalité », « X seul », « Aucun départ »)."""
    mini = f"MIN({plage})"
    nom_du_min = f"INDEX(${l0}$1:${lf}$1,MATCH({mini},{plage},0))"
    return (f'=IF(COUNT({plage})=0,"{AUCUNE}",IF(COUNT({plage})=1,{nom_du_min}&" seul",'
            f'IF(COUNTIF({plage},{mini})>1,"Égalité",{nom_du_min})))')


def construire_xlsx(produit, sites, contexte, large, details, releve_le, nuits, paire=None):
    """large : une ligne par date (colonnes date, mm, aaaa, jour, semaine, <site>, <site>__mp).
    details : lignes brutes filtrées (site, date, jour, semaine, prix_eur, meilleur_prix, compagnie)."""
    wb = Workbook()
    synth = wb.active
    synth.title = "Synthèse"
    comp = wb.create_sheet("Comparatif")
    det = wb.create_sheet("Détail")
    plusieurs = len(sites) > 1

    # ---------------------------------------------------------- Comparatif
    suffixe = f" ({paire[0]} / {paire[1]})" if paire else ""
    titres = ["Mois", "Semaine", "Jour", "Date", "Nombre de nuits"] + sites + \
             (["Moins cher", f"Écart{suffixe}"] if plusieurs else [])
    _entete(comp, 1, titres)
    col0 = 6                                          # première colonne de prix (F)
    col_fin = col0 + len(sites) - 1
    l0, lf = get_column_letter(col0), get_column_letter(col_fin)
    lignes = large.sort_values("date").reset_index(drop=True)
    for i, r in lignes.iterrows():
        n = i + 2
        _ecrire(comp, n, 1, f"{MOIS_FR[int(r['mm'])]} {int(r['aaaa'])}")
        _ecrire(comp, n, 2, f"Semaine {int(r['semaine'])}", centre=True)
        _ecrire(comp, n, 3, r["jour"], centre=True)
        _ecrire(comp, n, 4, r["date"].to_pydatetime(), FORMAT_DATE, centre=True)
        _ecrire(comp, n, 5, int(nuits), centre=True)
        for k, s in enumerate(sites):
            prix = r[s]
            manquant = prix != prix and r.get(f"{s}__echec") == True  # noqa: E712
            vide = NON_RELEVE if manquant else AUCUNE
            c = _ecrire(comp, n, col0 + k, vide if prix != prix else int(prix), FORMAT_EUR)
            if prix != prix:
                c.font = Font(name=POLICE, color="B45309" if manquant else "9AA5B1")
            if r.get(f"{s}__mp") == True:  # noqa: E712 (la valeur peut être NaN)
                c.comment = Comment("Meilleur prix du site", "Comparateur")
        if plusieurs:
            plage = f"{l0}{n}:{lf}{n}"
            # « Moins cher » : toujours le moins cher de la ligne, parmi tous les sites
            if any(r.get(f"{x}__echec") == True for x in sites):  # noqa: E712 (un site manque)
                _ecrire(comp, n, col_fin + 1, INCOMPLET, centre=True)
            else:
                _ecrire(comp, n, col_fin + 1, formule_moins_cher(plage, l0, lf), centre=True)
            # « Écart » : entre les deux sites cochés, sinon plus cher − moins cher
            if r.get("incomplet_ecart", r.get("incomplet")) == True:  # noqa: E712
                _ecrire(comp, n, col_fin + 2, INCOMPLET, centre=True)
            elif paire:
                a, b = (f"{get_column_letter(col0 + sites.index(s))}{n}" for s in paire)
                _ecrire(comp, n, col_fin + 2, f'=IF(AND(ISNUMBER({a}),ISNUMBER({b})),ABS({a}-{b}),'
                        f'IF(OR(ISNUMBER({a}),ISNUMBER({b})),"{UN_SEUL}","{AUCUNE}"))', FORMAT_EUR)
            else:
                _ecrire(comp, n, col_fin + 2, f'=IF(COUNT({plage})=0,"{AUCUNE}",IF(COUNT({plage})=1,"{UN_SEUL}",'
                        f'MAX({plage})-MIN({plage})))', FORMAT_EUR)
    derniere = len(lignes) + 1
    if plusieurs and derniere >= 2:
        comp.conditional_formatting.add(
            f"{l0}2:{lf}{derniere}",
            FormulaRule(formula=[f"AND(ISNUMBER({l0}2),{l0}2=MIN(${l0}2:${lf}2))"],
                        # dans une mise en forme conditionnelle, Excel lit la couleur dans bgColor
                        fill=PatternFill("solid", fgColor=VERT, bgColor=VERT), font=Font(name=POLICE, bold=True)))
    comp.freeze_panes = "F2"
    comp.auto_filter.ref = f"A1:{get_column_letter(len(titres))}{derniere}"
    for j, largeur in enumerate([16, 12, 8, 12, 10] + [14] * len(sites)
                                + ([18, max(10, len(suffixe) + 8)] if plusieurs else []), start=1):
        comp.column_dimensions[get_column_letter(j)].width = largeur

    # ---------------------------------------------------------- Synthèse
    synth["A1"] = "Comparatif de prix" if plusieurs else "Relevé de prix"
    synth["A1"].font = Font(name=POLICE, bold=True, size=16, color=BLEU_FONCE)
    infos = [f"Produit : {produit}", *contexte,
             f"Prix relevés le {releve_le:%d/%m/%Y} à {releve_le:%H:%M} — prix par personne (base chambre double) ; "
             "les tarifs évoluent en continu."]
    if plusieurs:
        infos.append("Onglet Comparatif : le prix le plus bas de chaque date est surligné en vert ; "
                     + (f"Écart = différence de prix entre {paire[0]} et {paire[1]}." if paire else
                        "Écart = prix le plus haut − prix le plus bas parmi les sites."))
    for i, texte in enumerate(infos, start=2):
        _ecrire(synth, i, 1, texte, gras=(i == 2))
    t = len(infos) + 3
    _entete(synth, t, ["Site", "Dates avec prix", "Moins cher (nb dates)", "Prix moyen", "Prix le plus bas",
                       "Date du prix le plus bas"])
    col_mc = get_column_letter(col_fin + 1)  # colonne « Moins cher » (tous les sites)
    for k, s in enumerate(sites):
        n, lettre = t + 1 + k, get_column_letter(col0 + k)
        plage = f"Comparatif!${lettre}$2:${lettre}${max(derniere, 2)}"
        _ecrire(synth, n, 1, s, gras=True)
        _ecrire(synth, n, 2, f"=COUNT({plage})", centre=True)
        _ecrire(synth, n, 3, f"=COUNTIF(Comparatif!${col_mc}$2:${col_mc}${max(derniere, 2)},A{n})" if plusieurs
                else "—", centre=True)
        _ecrire(synth, n, 4, f'=IFERROR(ROUND(AVERAGE({plage}),0),"{AUCUNE}")', FORMAT_EUR)
        _ecrire(synth, n, 5, f'=IF(COUNT({plage})=0,"{AUCUNE}",MIN({plage}))', FORMAT_EUR)
        _ecrire(synth, n, 6, f'=IFERROR(INDEX(Comparatif!$D$2:$D${max(derniere, 2)},MATCH(MIN({plage}),{plage},0)),"{AUCUNE}")',
                FORMAT_DATE, centre=True)
    if plusieurs:
        n = t + 1 + len(sites)
        _ecrire(synth, n, 1, "Égalités", gras=True)
        _ecrire(synth, n, 3, f'=COUNTIF(Comparatif!${col_mc}$2:${col_mc}${max(derniere, 2)},"Égalité")', centre=True)
        for j in range(1, 7):
            synth.cell(n, j).fill = PatternFill("solid", fgColor=GRIS)
    for j, largeur in enumerate([18, 16, 20, 14, 16, 22], start=1):
        synth.column_dimensions[get_column_letter(j)].width = largeur

    # ---------------------------------------------------------- Détail
    # Colonnes fournies par certains sites seulement (relevé navigateur Promoséjours) : ajoutées si présentes
    extras = [(col, titre) for col, titre in EXTRAS_DETAIL
              if col in details.columns and details[col].fillna("").astype(str).str.strip().ne("").any()]
    _entete(det, 1, ["Site", "Date", "Jour", "Semaine", "Nombre de nuits", "Prix", "Meilleur prix du site",
                     "Compagnie aérienne"] + [titre for _, titre in extras])
    d = details.sort_values(["date", "site"]).reset_index(drop=True)
    for i, r in d.iterrows():
        n = i + 2
        _ecrire(det, n, 1, r["site"])
        _ecrire(det, n, 2, r["date"].to_pydatetime(), FORMAT_DATE, centre=True)
        _ecrire(det, n, 3, r["jour"], centre=True)
        _ecrire(det, n, 4, f"Semaine {int(r['semaine'])}", centre=True)
        _ecrire(det, n, 5, int(nuits), centre=True)
        _ecrire(det, n, 6, int(r["prix_eur"]), FORMAT_EUR)
        _ecrire(det, n, 7, "Oui" if r["meilleur_prix"] else "", centre=True)
        _ecrire(det, n, 8, r.get("compagnie") or NON_PRECISEE)
        for k, (col, _) in enumerate(extras, start=9):
            v = r.get(col)
            if col == "retour" and isinstance(v, str) and v:
                _ecrire(det, n, k, datetime.strptime(v, "%Y-%m-%d"), FORMAT_DATE, centre=True)
            else:
                _ecrire(det, n, k, v if isinstance(v, str) and v else NON_PRECISEE, centre=True)
    derniere_col = get_column_letter(8 + len(extras))
    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:{derniere_col}{len(d) + 1}"
    for j, largeur in enumerate([16, 12, 8, 12, 10, 12, 20, 22] + [16] * len(extras), start=1):
        det.column_dimensions[get_column_letter(j)].width = largeur

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
