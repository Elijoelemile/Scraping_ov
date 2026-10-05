"""Comparatif Word Ovoyages / Exotismes : tarifs du mercredi et du samedi, semaine par semaine.

Les deux sites doivent porter sur le même hôtel, avec la même ville de départ et la même durée.

Exemples :
    python comparer_sites.py
    python comparer_sites.py --nuits 7 --debut 10-2026 --fin 08-2027
"""
import argparse
import calendar
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm

import scrape_exotismes as ex
import scrape_ovoyages as ov
from rapport_ovoyages import (MOIS_FR, BLEU_FONCE, GRIS, mois_entre, ombrer, ecrire, ecrire_ligne, ov_ville)

# Hôtel Coral Level at Iberostar Selection Bavaro 5*, vendu par les deux sites
OVOYAGES_URL = ov.BASE + "/voyages-rep-dominicaine/vacances-punta-cana/hotel-coral-level-at-iberostar-selection-bavaro-5-108504"
EXOTISMES_URL = ex.BASE + "/voyages/184-sejour-republique-dominicaine/25573-iberostar-selection-coral-bavaro.jsf"
VERT = "C6EFCE"        # site le moins cher
MERCREDI, SAMEDI = 2, 5


def prix_par_date(lignes):
    return {l["date"]: l for l in lignes if l.get("prix_eur")}


def cellule_prix(cell, info, moins_cher):
    if not info:
        ecrire(cell, "pas de départ", couleur="9AA5B1", taille=9)
        return
    ecrire(cell, f"{info['prix_eur']} €", gras=True, taille=11)
    if moins_cher:
        ombrer(cell, VERT)
    if info.get("meilleur_prix"):
        ecrire_ligne(cell, "Meilleur prix du site", taille=7)


def tableau_comparatif(doc, mm, aaaa, ovo, exo):
    entetes = ["Semaine", "Jour", "Date", "Ovoyages", "Exotismes", "Écart", "Moins cher"]
    largeurs = [Cm(2.2), Cm(2.0), Cm(1.8), Cm(2.8), Cm(2.8), Cm(2.4), Cm(2.6)]
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, entetes):
        ecrire(c, h, gras=True, couleur="FFFFFF")
        ombrer(c, BLEU_FONCE)

    stats = {"Ovoyages": 0, "Exotismes": 0, "Égalité": 0}
    ecarts = []
    semaines = calendar.Calendar(firstweekday=0).monthdatescalendar(aaaa, mm)
    for num, jours in enumerate(semaines, start=1):
        lignes_semaine = []
        for idx, nom in ((MERCREDI, "Mercredi"), (SAMEDI, "Samedi")):
            d = jours[idx]
            cells = t.add_row().cells
            lignes_semaine.append(cells)
            hors_mois = d.month != mm
            gris = "9AA5B1" if hors_mois else None
            ecrire(cells[1], nom, couleur=gris)
            ecrire(cells[2], f"{d.day:02d}/{d.month:02d}", couleur=gris)
            if hors_mois:
                for c in cells[1:]:
                    ombrer(c, GRIS)
                ecrire(cells[3], "hors mois", couleur=gris, taille=9)
                continue
            o, e = ovo.get(d.isoformat()), exo.get(d.isoformat())
            po, pe = (o or {}).get("prix_eur"), (e or {}).get("prix_eur")
            cellule_prix(cells[3], o, bool(po and pe and po < pe))
            cellule_prix(cells[4], e, bool(po and pe and pe < po))
            if po and pe:
                ecart = pe - po
                ecarts.append(ecart)
                ecrire(cells[5], f"{ecart:+d} €" if ecart else "0 €", taille=10)
                gagnant = "Égalité" if ecart == 0 else ("Ovoyages" if ecart > 0 else "Exotismes")
                stats[gagnant] += 1
                ecrire(cells[6], gagnant, gras=gagnant != "Égalité")
            else:
                seul = "Ovoyages" if po else "Exotismes" if pe else None
                ecrire(cells[5], "—", couleur="9AA5B1")
                ecrire(cells[6], f"{seul} seul" if seul else "—", couleur="9AA5B1", taille=9)
        # « Semaine N » fusionnée sur ses deux lignes (mercredi + samedi)
        fusion = lignes_semaine[0][0].merge(lignes_semaine[1][0])
        ecrire(fusion, f"Semaine {num}", gras=True)

    for i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, w in zip(row.cells, largeurs):
            c.width = w
            if i < len(t.rows) - 1:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    return stats, ecarts


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--ville", default="PAR", help="code ville de départ (défaut : PAR)")
    p.add_argument("--nuits", type=int, default=5, help="nombre de nuits (défaut : 5)")
    p.add_argument("--debut", type=ov.parse_mois, default=(10, 2026), help="premier mois MM-AAAA (défaut 10-2026)")
    p.add_argument("--fin", type=ov.parse_mois, default=(8, 2027), help="dernier mois MM-AAAA (défaut 08-2027)")
    p.add_argument("--pause", type=float, default=3.0, help="secondes entre deux requêtes (défaut : 3)")
    p.add_argument("--sortie", default="comparatif_ovoyages_exotismes.docx", help="fichier Word de sortie")
    args = p.parse_args()
    ov.PAUSE = ex.PAUSE = args.pause
    ville = args.ville.upper()

    code_ovo = ov.code_catalogue(OVOYAGES_URL)
    if args.nuits not in ov.durees(code_ovo, ville):
        raise SystemExit(f"Ovoyages ne propose pas {args.nuits} nuits au départ de {ville}")
    grille = ex.GrilleExotismes(ex.trouver_grille_id(EXOTISMES_URL))
    grille.change("nbNuits", str(args.nuits))
    if ville != "PAR":
        grille.change("villeDepart", ville)
    mois_exo = {v for v, _ in grille.options("dateDepart")}

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("Comparatif Ovoyages / Exotismes", level=0)
    infos = doc.add_paragraph()
    infos.add_run("Hôtel Coral Level at Iberostar Selection Bávaro 5*").bold = True
    infos.add_run(f"\nTarifs du mercredi et du samedi · Départ : {ov_ville(ville)} · Durée : {args.nuits} nuits · "
                  f"{MOIS_FR[args.debut[0]]} {args.debut[1]} à {MOIS_FR[args.fin[0]]} {args.fin[1]}")
    infos.add_run(f"\nPrix par personne (base chambre double), relevés le {datetime.now():%d/%m/%Y} à {datetime.now():%H:%M}. "
                  "Les tarifs évoluent en continu.")
    infos.add_run(f"\nSources : {OVOYAGES_URL}\n{EXOTISMES_URL}")
    note = doc.add_paragraph()
    note.add_run("Écart = prix Exotismes − prix Ovoyages : positif, Ovoyages est moins cher ; négatif, Exotismes "
                 "est moins cher. Le prix le plus bas est surligné en vert. Une semaine correspond à une ligne du "
                 "calendrier des sites (du lundi au dimanche).").italic = True

    # Tableau de synthèse rempli après la collecte
    doc.add_heading("Synthèse", level=1)
    synthese = doc.add_table(rows=1, cols=5)
    synthese.style = "Table Grid"
    synthese.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(synthese.rows[0].cells, ["Mois", "Ovoyages moins cher", "Exotismes moins cher",
                                              "Égalité", "Écart moyen"]):
        ecrire(c, h, gras=True, couleur="FFFFFF")
        ombrer(c, BLEU_FONCE)
    total = {"Ovoyages": 0, "Exotismes": 0, "Égalité": 0}
    tous_ecarts = []

    for mm, aaaa in mois_entre(args.debut, args.fin):
        ovo = prix_par_date(ov.calendrier(code_ovo, ville, args.nuits, mm, aaaa)[0])
        cle = f"{mm:02d}-{aaaa}"
        exo = {}
        if cle in mois_exo:
            grille.change("dateDepart", cle, render="pageformulaire infoAutreBrochure")
            exo = prix_par_date(grille.calendrier())
        print(f"{cle} : Ovoyages {len(ovo)} dates, Exotismes {len(exo)} dates")

        if (mm, aaaa) == args.debut:
            doc.add_page_break()  # la synthèse reste seule en première page
        else:
            doc.add_paragraph()
        h = doc.add_heading(f"{MOIS_FR[mm]} {aaaa}", level=1)
        h.paragraph_format.keep_with_next = True
        stats, ecarts = tableau_comparatif(doc, mm, aaaa, ovo, exo)

        cells = synthese.add_row().cells
        ecrire(cells[0], f"{MOIS_FR[mm]} {aaaa}", gras=True)
        for c, k in zip(cells[1:4], ["Ovoyages", "Exotismes", "Égalité"]):
            ecrire(c, str(stats[k]))
            total[k] += stats[k]
        ecrire(cells[4], f"{round(sum(ecarts) / len(ecarts)):+d} €" if ecarts else "—")
        tous_ecarts += ecarts

    cells = synthese.add_row().cells
    ecrire(cells[0], "Total", gras=True)
    for c, k in zip(cells[1:4], ["Ovoyages", "Exotismes", "Égalité"]):
        ecrire(c, str(total[k]), gras=True)
    ecrire(cells[4], f"{round(sum(tous_ecarts) / len(tous_ecarts)):+d} €" if tous_ecarts else "—", gras=True)
    for c in cells:
        ombrer(c, GRIS)
    for row in synthese.rows:
        for c, w in zip(row.cells, [Cm(3.4), Cm(3.4), Cm(3.4), Cm(2.4), Cm(3.0)]):
            c.width = w

    doc.save(args.sortie)
    print(f"\nComparatif écrit dans {args.sortie}")


if __name__ == "__main__":
    main()
